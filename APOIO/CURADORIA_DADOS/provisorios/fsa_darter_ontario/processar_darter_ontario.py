from __future__ import annotations

import csv
import json
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "provisorios" / "fsa_darter_ontario" / "DarterOnt.csv"
MAIN = ROOT / "dados_organizados_limpos_para_pacote_aulas.xlsx"
FALLBACK_MAIN = ROOT / "dados_organizados_limpos_para_pacote_aulas_com_darter_ontario.xlsx"
REGISTRY = ROOT / "REGISTRO_FONTES_DADOS_EAPA.csv"
OUT_CSV = ROOT / "provisorios" / "fsa_darter_ontario" / "darter_ontario_limpo.csv"
VALIDATION_JSON = ROOT / "provisorios" / "fsa_darter_ontario" / "validacao_darter_ontario.json"

CLEAN_SHEET = "darter_ontario1"
DATASET_ID = "darter_ontario_fsa_reid_2004"


def rankdata_average(values):
    indexed = sorted(enumerate(values), key=lambda item: item[1])
    ranks = [0.0] * len(values)
    i = 0
    while i < len(indexed):
        j = i
        while j + 1 < len(indexed) and indexed[j + 1][1] == indexed[i][1]:
            j += 1
        avg_rank = (i + 1 + j + 1) / 2
        for k in range(i, j + 1):
            ranks[indexed[k][0]] = avg_rank
        i = j + 1
    return ranks


def mann_whitney_summary(df):
    values = df["idade_anos"].tolist()
    groups = df["rio"].tolist()
    ranks = rankdata_average(values)
    n_salmon = sum(1 for g in groups if g == "Salmon")
    n_trent = sum(1 for g in groups if g == "Trent")
    rank_salmon = sum(r for r, g in zip(ranks, groups) if g == "Salmon")
    u_salmon = rank_salmon - n_salmon * (n_salmon + 1) / 2
    u_trent = n_salmon * n_trent - u_salmon
    ties = pd.Series(values).value_counts()
    return {
        "n_salmon": n_salmon,
        "n_trent": n_trent,
        "mediana_idade_salmon": float(df.loc[df["rio"] == "Salmon", "idade_anos"].median()),
        "mediana_idade_trent": float(df.loc[df["rio"] == "Trent", "idade_anos"].median()),
        "u_salmon": float(u_salmon),
        "u_trent": float(u_trent),
        "u_menor": float(min(u_salmon, u_trent)),
        "idades_com_empate": {str(int(k)): int(v) for k, v in ties.items() if v > 1},
        "ha_empates": bool((ties > 1).any()),
    }


def clean_data():
    raw = pd.read_csv(RAW)
    out = pd.DataFrame(
        {
            "id_peixe": [f"DO{i:02d}" for i in range(1, len(raw) + 1)],
            "idade_anos": raw["age"].astype(int),
            "comprimento_total_mm": pd.to_numeric(raw["tl"], errors="coerce"),
            "rio": raw["river"].astype(str),
        }
    )
    return raw, out


def validate(raw, out):
    mw = mann_whitney_summary(out)
    summary = {
        "linhas_originais": int(raw.shape[0]),
        "colunas_originais": int(raw.shape[1]),
        "linhas_limpo": int(out.shape[0]),
        "colunas_limpo": int(out.shape[1]),
        "rios": out["rio"].value_counts().to_dict(),
        "idades": out["idade_anos"].value_counts().sort_index().to_dict(),
        "ausentes_por_coluna": out.isna().sum().to_dict(),
        "duplicatas_id_peixe": int(out["id_peixe"].duplicated().sum()),
        "duplicatas_completas": int(out.duplicated().sum()),
        "mann_whitney": mw,
    }
    VALIDATION_JSON.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def autosize(ws, max_width=48):
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        values = [ws.cell(row, col_idx).value for row in range(1, min(ws.max_row, 80) + 1)]
        width = min(max_width, max(10, max(len(str(v)) if v is not None else 0 for v in values) + 2))
        ws.column_dimensions[letter].width = width


def set_header_comments(ws, comments):
    for cell in ws[1]:
        if cell.value in comments:
            cell.comment = Comment(comments[cell.value], "EAPA")


def remove_sheet_if_exists(wb, sheet_name):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]


def update_workbook(out, summary):
    wb = openpyxl.load_workbook(MAIN)
    remove_sheet_if_exists(wb, CLEAN_SHEET)
    ws = wb.create_sheet(CLEAN_SHEET)
    ws.append(list(out.columns))
    for record in out.to_dict(orient="records"):
        ws.append([record[col] for col in out.columns])

    header_fill = PatternFill("solid", fgColor="314E7A")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)

    set_header_comments(
        ws,
        {
            "id_peixe": "Identificador sequencial criado na curadoria; cada linha representa um peixe capturado.",
            "idade_anos": "Idade estimada por leitura de otolitos, em anos. Variavel resposta sugerida para Mann-Whitney.",
            "comprimento_total_mm": "Comprimento total do peixe, em milimetros.",
            "rio": "Rio de captura; grupos independentes Salmon e Trent.",
        },
    )

    info = wb["info_conjuntos"]
    existing_ids = {info.cell(row, 1).value: row for row in range(2, info.max_row + 1)}
    row_idx = existing_ids.get(DATASET_ID, info.max_row + 1)
    info_values = {
        "id_conjunto": DATASET_ID,
        "nome_dataset_pacote": "darter_ontario",
        "nome_curto": "Idade e comprimento da percina-do-canal em dois rios de Ontario",
        "objetivo_didatico": "Comparar a distribuicao das idades de Percina copelandi entre dois rios independentes usando o teste de Mann-Whitney.",
        "area": "bioecologia pesqueira; biologia de peixes",
        "tema": "idade; comprimento; comparacao de dois grupos; percina-do-canal",
        "analise_principal": "Mann-Whitney",
        "modelo_sugerido": "idade_anos ~ rio",
        "modelo_sugerido_r": "wilcox.test(idade_anos ~ rio, data = darter_ontario, exact = FALSE)",
        "unidade_observacional": "peixe individual",
        "status_limpeza": "limpo_validado",
        "tipo_adaptacao": "variaveis traduzidas para portugues; dados mantidos no nivel individual",
        "aba_limpa": CLEAN_SHEET,
        "linhas_limpo": summary["linhas_limpo"],
        "colunas_limpo": summary["colunas_limpo"],
        "linhas_dados_originais": summary["linhas_originais"],
        "colunas_dados_originais": summary["colunas_originais"],
        "fonte": "FSAdata::DarterOnt, reconstruido a partir da Figura 2 de Reid (2004), com dados de Percina copelandi dos rios Salmon e Trent, Ontario.",
        "doi": "10.1080/02705060.2004.9664917",
        "url_documentacao": "https://fishr-core-team.github.io/FSAdata/reference/DarterOnt.html",
        "licenca": "FSAdata GPL-2 ou GPL-3",
        "uso_recomendado": "atividade_consolidacao_aula_livro_pacote",
        "arquivo_bruto_local": "provisorios/fsa_darter_ontario/DarterOnt.csv",
        "observacoes": "Base pequena para Mann-Whitney. Idade e uma variavel discreta com muitos empates; usar aproximacao com correcao para empates, nao teste exato. Evitar interpretar apenas como comparacao de medias.",
    }
    headers = [info.cell(1, col).value for col in range(1, info.max_column + 1)]
    for col_idx, header in enumerate(headers, start=1):
        info.cell(row_idx, col_idx).value = info_values.get(header)
    autosize(info)
    try:
        wb.save(MAIN)
        return str(MAIN)
    except PermissionError:
        wb.save(FALLBACK_MAIN)
        return str(FALLBACK_MAIN)


def update_registry(summary):
    registry_row = {
        "id_fonte": "darter_ontario_fsa_reid_2004",
        "origem": "FSAdata::DarterOnt. Idades e comprimentos de Percina copelandi dos rios Salmon e Trent, Ontario; reconstruido da Figura 2 de Reid (2004). Artigo DOI: 10.1080/02705060.2004.9664917.",
        "tipo_fonte": "dataset_pacote_r",
        "data_acesso": "2026-07-11",
        "licenca": "GPL-2 ou GPL-3",
        "status_licenca": "publico_aberto",
        "area": "bioecologia pesqueira; biologia de peixes",
        "tema": "idade; comprimento; comparacao de dois grupos; Mann-Whitney",
        "unidade_observacional": "peixe individual",
        "estrutura_estatistica": "Mann-Whitney; graficos de distribuicao; regressao idade-comprimento",
        "destino_pretendido": "atividade_consolidacao_aula_livro_pacote",
        "decisao": "aceitar",
        "arquivo_local": "dados_organizados_limpos_para_pacote_aulas.xlsx; provisorios/fsa_darter_ontario/DarterOnt.csv",
        "observacoes": f"Base limpa adicionada na aba {CLEAN_SHEET} com {summary['linhas_limpo']} peixes e {summary['colunas_limpo']} variaveis. Ha empates em idade; usar wilcox.test(..., exact = FALSE).",
    }
    with REGISTRY.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        rows = [row for row in reader if row.get("id_fonte") != registry_row["id_fonte"]]
        fieldnames = reader.fieldnames
    rows.append(registry_row)
    with REGISTRY.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main():
    raw, out = clean_data()
    summary = validate(raw, out)
    if summary["linhas_limpo"] != 54 or set(summary["rios"]) != {"Salmon", "Trent"}:
        raise SystemExit("Falha na validacao de dimensoes ou grupos.")
    out.to_csv(OUT_CSV, index=False, encoding="utf-8")
    saved_workbook = update_workbook(out, summary)
    update_registry(summary)
    summary["workbook_salvo"] = saved_workbook
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
