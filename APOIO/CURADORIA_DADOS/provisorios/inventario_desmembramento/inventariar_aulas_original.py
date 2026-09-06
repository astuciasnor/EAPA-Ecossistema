from pathlib import Path
import json
from openpyxl import load_workbook

source = Path(r"D:\Claude\eapa\EAPADados\CURADORIA_DADOS\fontes_containers_originais\aulas_bioestatistica_ORIGINAL.xlsx")
output = Path("inventario_aulas_original.json")

workbook = load_workbook(source, read_only=True, data_only=False)
inventory = []
for index, sheet in enumerate(workbook.worksheets, start=1):
    max_row = sheet.max_row or 1
    max_column = sheet.max_column or 1
    preview = []
    for row in sheet.iter_rows(min_row=1, max_row=min(max_row, 8), max_col=min(max_column, 12), values_only=True):
        preview.append(["" if value is None else str(value).strip() for value in row])
    inventory.append({
        "sheet_index": index,
        "sheet_name": sheet.title,
        "rows": max_row,
        "columns": max_column,
        "preview": preview,
    })

output.write_text(json.dumps(inventory, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(inventory, ensure_ascii=False, indent=2))
