from __future__ import annotations

import csv
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "provisorios" / "ropke_2025_bagres_madeira_mendeley" / "diet of large catfishes from Madeira River_mendeley data.csv"
MAIN = ROOT / "dados_organizados_limpos_para_pacote_aulas.xlsx"
REGISTRY = ROOT / "REGISTRO_FONTES_DADOS_EAPA.csv"
OUT_CSV = ROOT / "provisorios" / "ropke_2025_bagres_madeira_mendeley" / "bagres_dieta_madeira_limpo.csv"
VALIDATION_JSON = ROOT / "provisorios" / "ropke_2025_bagres_madeira_mendeley" / "validacao_bagres_dieta_madeira.json"

CLEAN_SHEET = "bagres_dieta_madeira1"
DATASET_ID = "bagres_dieta_madeira_ropke_2025"


MONTH_PT = {
    "January": "janeiro",
    "February": "fevereiro",
    "March": "marco",
    "April": "abril",
    "May": "maio",
    "June": "junho",
    "July": "julho",
    "August": "agosto",
    "September": "setembro",
    "October": "outubro",
    "November": "novembro",
    "December": "dezembro",
}

SEASON_PT = {
    "Low-water": "vazante_seca",
    "High-water": "cheia",
}

POPULAR_FIX = {
    "Babão": "babao",
    "Filhote": "filhote",
    "Dourada": "dourada",
    "Piramutaba": "piramutaba",
    "Surubim": "surubim",
    "Jau": "jau",
    "Barbachata": "barbachata",
    "Caparari": "caparari",
}

FISH_NI = {"fish n.i."}
FISH_REMAINS = {"bones", "scales"}
FISH_ORDERS = {"Characiformes", "Clupeiformes", "Gymnotyformes", "Siluriformes"}
FISH_FAMILIES = {
    "Anostomidae",
    "Characidae",
    "Curimatidae",
    "Doradidae",
    "Hemiodontidae",
    "Pimelodidae",
    "Myleinae",
    "Tetragonopterinae",
}
INVERTEBRATES = {
    "Coleoptera",
    "Megaloptera",
    "Trichoptera",
    "Trichodactylus sp.",
    "Macrobrachium spp.",
    "reamins of invertebrates",
}


def strip_accents(text):
    return "".join(
        ch for ch in unicodedata.normalize("NFKD", str(text)) if not unicodedata.combining(ch)
    )


def snake(text):
    text = strip_accents(str(text).strip().lower())
    text = text.replace("n.i.", "ni").replace("spp.", "spp").replace("sp.", "sp")
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return re.sub(r"_+", "_", text).strip("_")


def norm_text(value):
    if pd.isna(value):
        return None
    text = str(value).strip()
    return text if text else None


def norm_locality(value):
    text = norm_text(value)
    if not text:
        return None
    text = re.sub(r"\s+", " ", text).strip()
    fixes = {
        "Cachoeira do macaco": "Cachoeira do Macaco",
        "jaciparaná": "Jaciparaná",
        "Jaciparaná ": "Jaciparaná",
    }
    return fixes.get(text, text)


def to_number(series):
    return pd.to_numeric(series.astype(str).str.replace(",", ".", regex=False), errors="coerce")


def classify_diet_cols(diet_cols):
    grouped = {
        "volume_peixe_nao_identificado_pct": [],
        "volume_restos_peixe_pct": [],
        "volume_ordens_peixes_pct": [],
        "volume_familias_peixes_pct": [],
        "volume_peixes_identificados_pct": [],
        "volume_invertebrados_pct": [],
    }
    for col in diet_cols:
        if col in FISH_NI:
            grouped["volume_peixe_nao_identificado_pct"].append(col)
        elif col in FISH_REMAINS:
            grouped["volume_restos_peixe_pct"].append(col)
        elif col in FISH_ORDERS:
            grouped["volume_ordens_peixes_pct"].append(col)
        elif col in FISH_FAMILIES:
            grouped["volume_familias_peixes_pct"].append(col)
        elif col in INVERTEBRATES:
            grouped["volume_invertebrados_pct"].append(col)
        else:
            grouped["volume_peixes_identificados_pct"].append(col)
    return grouped


def dominant_item(row, diet_cols):
    values = row[diet_cols].fillna(0)
    max_value = values.max()
    if max_value <= 0:
        return None, None
    return values.idxmax(), float(max_value)


def read_and_clean():
    df = pd.read_csv(RAW, sep=";", encoding="mac_roman")
    df.columns = [c.strip() for c in df.columns]
    diet_cols = list(df.columns[10:])
    diet_numeric = df[diet_cols].apply(to_number).fillna(0)
    original_totals = diet_numeric.sum(axis=1)
    diet_normalized = diet_numeric.div(original_totals.replace(0, pd.NA), axis=0).fillna(0) * 100

    groups = classify_diet_cols(diet_cols)
    out = pd.DataFrame()
    out["id_estomago"] = [f"BM{i:03d}" for i in range(1, len(df) + 1)]
    out["nome_popular"] = df["popular_name"].map(lambda x: POPULAR_FIX.get(str(x).strip(), snake(x)))
    out["especie_predador"] = df["specie"].map(norm_text)
    out["genero_predador"] = out["especie_predador"].str.split().str[0]
    out["localidade"] = df["localilty"].map(norm_locality)
    out["trecho_rio"] = df["area"].map(norm_text)
    out["periodo_hidrologico"] = df["season"].map(SEASON_PT).fillna(df["season"].map(snake))
    out["mes"] = df["month"].map(MONTH_PT).fillna(df["month"].map(snake))
    out["ano"] = to_number(df["year"]).astype("Int64")
    out["comprimento_padrao_cm"] = to_number(df["standard_length (cm)"])
    out["peso_g"] = to_number(df["weigth (g)"])
    out["grau_replecao"] = to_number(df["stomach_fullness"]).astype("Int64")
    out["volume_total_original_pct"] = original_totals.round(4)
    out["volume_reescalonado_100"] = original_totals.round(6) != 100

    for new_col, cols in groups.items():
        out[new_col] = diet_normalized[cols].sum(axis=1).round(4)

    out["volume_total_dieta_pct"] = diet_normalized.sum(axis=1).round(4)
    out["riqueza_itens_alimentares"] = (diet_numeric > 0).sum(axis=1).astype(int)
    dominant = diet_numeric.apply(lambda row: dominant_item(row, diet_cols), axis=1)
    out["item_dominante_original"] = [item for item, _ in dominant]
    out["volume_item_dominante_pct"] = [value for _, value in dominant]
    out["dieta_mista"] = out["riqueza_itens_alimentares"] > 1

    return out, df, diet_numeric, diet_normalized, groups


def validate(out, raw, diet_numeric, diet_normalized, groups):
    totals = out["volume_total_dieta_pct"]
    original_totals = out["volume_total_original_pct"]
    summary = {
        "linhas_originais": int(raw.shape[0]),
        "colunas_originais": int(raw.shape[1]),
        "linhas_limpo": int(out.shape[0]),
        "colunas_limpo": int(out.shape[1]),
        "n_especies": int(out["especie_predador"].nunique()),
        "especies": out["especie_predador"].value_counts().to_dict(),
        "periodos_hidrologicos": out["periodo_hidrologico"].value_counts().to_dict(),
        "localidades": int(out["localidade"].nunique()),
        "itens_alimentares_originais": int(diet_numeric.shape[1]),
        "itens_alimentares_ativos": int((diet_numeric.sum(axis=0) > 0).sum()),
        "agrupamento_itens": {k: len(v) for k, v in groups.items()},
        "volume_total_min": float(totals.min()),
        "volume_total_max": float(totals.max()),
        "volume_total_original_min": float(original_totals.min()),
        "volume_total_original_max": float(original_totals.max()),
        "linhas_volume_total_original_diferente_100": int((original_totals.round(6) != 100).sum()),
        "linhas_reescalonadas": int(out["volume_reescalonado_100"].sum()),
        "linhas_volume_total_100_apos_reescala": int((totals.round(6) == 100).sum()),
        "linhas_volume_total_diferente_100_apos_reescala": int((totals.round(6) != 100).sum()),
        "ausentes_por_coluna": out.isna().sum().to_dict(),
        "duplicatas_id_estomago": int(out["id_estomago"].duplicated().sum()),
        "duplicatas_completas": int(out.duplicated().sum()),
    }
    VALIDATION_JSON.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def write_clean_csv(out):
    out.to_csv(OUT_CSV, index=False, encoding="utf-8")


def remove_sheet_if_exists(wb, sheet_name):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]


def autosize(ws, max_width=46):
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        values = [ws.cell(row, col_idx).value for row in range(1, min(ws.max_row, 80) + 1)]
        width = min(max_width, max(10, max(len(str(v)) if v is not None else 0 for v in values) + 2))
        ws.column_dimensions[letter].width = width


def set_header_comments(ws, comments):
    for cell in ws[1]:
        if cell.value in comments:
            cell.comment = Comment(comments[cell.value], "EAPA")


def update_workbook(out, summary):
    wb = openpyxl.load_workbook(MAIN)
    remove_sheet_if_exists(wb, CLEAN_SHEET)
    ws = wb.create_sheet(CLEAN_SHEET)
    ws.append(list(out.columns))
    for record in out.where(pd.notna(out), None).to_dict(orient="records"):
        ws.append([record[col] for col in out.columns])

    header_fill = PatternFill("solid", fgColor="6E3B1F")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)

    comments = {
        "id_estomago": "Identificador sequencial criado na curadoria; cada linha representa um estomago com conteudo analisado.",
        "nome_popular": "Nome popular do predador, padronizado em snake_case sem acento.",
        "especie_predador": "Especie de grande bagre/pimelodideo analisada.",
        "genero_predador": "Genero taxonomico do predador.",
        "localidade": "Localidade de desembarque/coleta conforme a fonte, com padronizacao minima de grafia.",
        "trecho_rio": "Trecho do rio Madeira/Mamore descrito na fonte.",
        "periodo_hidrologico": "Periodo hidrologico traduzido: cheia ou vazante_seca.",
        "mes": "Mes da coleta/desembarque em portugues, sem acento.",
        "ano": "Ano da coleta/desembarque.",
        "comprimento_padrao_cm": "Comprimento padrao do peixe predador, em centimetros.",
        "peso_g": "Peso total do peixe predador, em gramas.",
        "grau_replecao": "Grau visual de replecao estomacal: 1 <25%, 2 = 25-75%, 3 >75%.",
        "volume_total_original_pct": "Soma dos volumes dos itens alimentares antes da curadoria; valores diferentes de 100 indicam inconsistencia/arrendondamento no bruto.",
        "volume_reescalonado_100": "TRUE quando os itens alimentares foram reescalonados proporcionalmente para soma 100%.",
        "volume_peixe_nao_identificado_pct": "Volume relativo de peixe nao identificado no conteudo estomacal.",
        "volume_restos_peixe_pct": "Volume relativo de ossos e escamas.",
        "volume_ordens_peixes_pct": "Volume relativo de presas identificadas ate ordem.",
        "volume_familias_peixes_pct": "Volume relativo de presas identificadas ate familia ou subfamilia.",
        "volume_peixes_identificados_pct": "Volume relativo de presas identificadas ate genero, especie ou spp.",
        "volume_invertebrados_pct": "Volume relativo de insetos, caranguejos, camaroes e restos de invertebrados.",
        "volume_total_dieta_pct": "Soma dos volumes relativos dos itens alimentares; esperado = 100 em cada estomago.",
        "riqueza_itens_alimentares": "Numero de itens alimentares com volume maior que zero no registro original.",
        "item_dominante_original": "Item alimentar original com maior volume relativo no estomago.",
        "volume_item_dominante_pct": "Volume relativo do item alimentar dominante.",
        "dieta_mista": "TRUE quando mais de um item alimentar foi registrado no mesmo estomago.",
    }
    set_header_comments(ws, comments)

    info = wb["info_conjuntos"]
    existing_ids = {info.cell(row, 1).value: row for row in range(2, info.max_row + 1)}
    row_idx = existing_ids.get(DATASET_ID, info.max_row + 1)
    info_values = {
        "id_conjunto": DATASET_ID,
        "nome_dataset_pacote": "bagres_dieta_madeira",
        "nome_curto": "Dieta de grandes bagres do rio Madeira",
        "objetivo_didatico": "Atividade de consolidacao sobre dieta, sazonalidade hidrologica e particao de recursos em grandes bagres amazonicos.",
        "area": "bioecologia pesqueira; ecologia trofica; Amazonia",
        "tema": "dieta; bagres migradores; periodo hidrologico; particao de nicho; rio Madeira",
        "analise_principal": "qui-quadrado; ANOVA/GLM; ACP/AAH exploratoria",
        "modelo_sugerido": "volume_peixes_identificados_pct ~ especie_predador + periodo_hidrologico",
        "modelo_sugerido_r": "volume_peixes_identificados_pct ~ especie_predador + periodo_hidrologico",
        "unidade_observacional": "estomago de peixe predador com conteudo alimentar",
        "status_limpeza": "limpo_validado",
        "tipo_adaptacao": "dados individuais preservados; 85 itens alimentares originais agrupados em seis grupos troficos didaticos; volumes reescalonados para soma 100 quando necessario",
        "aba_limpa": CLEAN_SHEET,
        "linhas_limpo": summary["linhas_limpo"],
        "colunas_limpo": summary["colunas_limpo"],
        "linhas_dados_originais": summary["linhas_originais"],
        "colunas_dados_originais": summary["colunas_originais"],
        "fonte": "Ropke, C.; Zuanon, J. A. S.; Fonseca, M.; Lima, M. A.; Sant'Anna, I.; Gunther, H.; Torrente-Vilara, G.; Doria, C. (2025). Data for: Diet of large Catfishes from Madeira River. Mendeley Data.",
        "doi": "10.17632/b7kwsyzmyf.1",
        "url_documentacao": "https://data.mendeley.com/datasets/b7kwsyzmyf/1",
        "licenca": "CC BY 4.0",
        "uso_recomendado": "atividade_consolidacao_aula_livro_pacote",
        "arquivo_bruto_local": "provisorios/ropke_2025_bagres_madeira_mendeley/diet of large catfishes from Madeira River_mendeley data.csv",
        "observacoes": "Base lida com encoding MacRoman. Cabecalhos corrigidos apenas na documentacao/nomes curados. Tres estomagos tinham soma original diferente de 100% (99,99; 180; 478); os volumes foram reescalonados proporcionalmente e a soma original ficou registrada. Os itens detalhados foram agrupados para reduzir complexidade da atividade.",
    }
    headers = [info.cell(1, col).value for col in range(1, info.max_column + 1)]
    for col_idx, header in enumerate(headers, start=1):
        info.cell(row_idx, col_idx).value = info_values.get(header)
    autosize(info)

    wb.save(MAIN)


def update_registry(summary):
    registry_row = {
        "id_fonte": "ropke_2025_bagres_madeira_mendeley",
        "origem": "Ropke, C.; Zuanon, J. A. S.; Fonseca, M.; Lima, M. A.; Sant'Anna, I.; Gunther, H.; Torrente-Vilara, G.; Doria, C. (2025). Data for: Diet of large Catfishes from Madeira River. Mendeley Data. DOI: 10.17632/b7kwsyzmyf.1.",
        "tipo_fonte": "dataset_mendeley",
        "data_acesso": "2026-07-11",
        "licenca": "CC BY 4.0",
        "status_licenca": "publico_aberto",
        "area": "bioecologia pesqueira; ecologia trofica; Amazonia",
        "tema": "dieta; bagres migradores; periodo hidrologico; particao de nicho; rio Madeira",
        "unidade_observacional": "estomago de peixe predador",
        "estrutura_estatistica": "qui-quadrado; ANOVA/GLM; ACP; AAH; comparacao entre especies e periodo hidrologico",
        "destino_pretendido": "atividade_consolidacao_aula_livro_pacote",
        "decisao": "aceitar",
        "arquivo_local": "dados_organizados_limpos_para_pacote_aulas.xlsx; provisorios/ropke_2025_bagres_madeira_mendeley/diet of large catfishes from Madeira River_mendeley data.csv",
        "observacoes": f"Base limpa adicionada na aba {CLEAN_SHEET} com {summary['linhas_limpo']} estomagos e {summary['colunas_limpo']} variaveis. Os 85 itens alimentares originais foram agrupados em seis grupos troficos didaticos; {summary['linhas_reescalonadas']} linhas foram reescalonadas proporcionalmente para soma 100%.",
    }
    with REGISTRY.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        rows = [row for row in reader if row.get("id_fonte") != registry_row["id_fonte"]]
        fieldnames = reader.fieldnames
    rows.append(registry_row)
    with REGISTRY.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main():
    out, raw, diet_numeric, diet_normalized, groups = read_and_clean()
    summary = validate(out, raw, diet_numeric, diet_normalized, groups)
    if summary["linhas_volume_total_diferente_100_apos_reescala"] != 0:
        raise SystemExit("Falha: ha linhas cujo volume alimentar nao soma 100%. Conferir validacao.")
    write_clean_csv(out)
    update_workbook(out, summary)
    update_registry(summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
