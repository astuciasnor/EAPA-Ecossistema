from __future__ import annotations

import csv
import json
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from statistics import mean

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "provisorios" / "malyenge_2025_lobster_mendeley" / "Lobster Log sheet.xlsx"
MAIN = ROOT / "dados_organizados_limpos_para_pacote_aulas.xlsx"
REGISTRY = ROOT / "REGISTRO_FONTES_DADOS_EAPA.csv"
OUT_CSV = ROOT / "provisorios" / "malyenge_2025_lobster_mendeley" / "lagosta_namibia_limpo.csv"
VALIDATION_JSON = ROOT / "provisorios" / "malyenge_2025_lobster_mendeley" / "validacao_lagosta_namibia.json"


MONTH_SHEETS = {
    "Sept 2024": (2024, 9, "setembro", date(2024, 9, 28)),
    "Oct 2024": (2024, 10, "outubro", date(2024, 10, 23)),
    "Dec 2024": (2024, 12, "dezembro", date(2024, 12, 6)),
    "Jan 2025": (2025, 1, "janeiro", date(2025, 1, 29)),
    "Mar 2025": (2025, 3, "marco", date(2025, 3, 28)),
    "April 2025": (2025, 4, "abril", date(2025, 4, 17)),
}

COMBINED_MONTH = {
    "September": "setembro",
    "October": "outubro",
    "December": "dezembro",
    "January": "janeiro",
    "March": "marco",
    "April": "abril",
}

CLEAN_SHEET = "lagosta_namibia1"
DATASET_ID = "lagosta_namibia_malyenge_2025"


def as_float(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip().replace(",", "."))
    except ValueError:
        return None


def as_int(value):
    number = as_float(value)
    return None if number is None else int(number)


def norm_text(value):
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def norm_site(value):
    text = norm_text(value)
    if not text:
        return None
    text = text.upper().replace(" ", "")
    if "DIAZ" in text:
        return "DIAZ"
    if "SWB" in text:
        return "SWB"
    return text


def norm_sex(value):
    text = norm_text(value)
    if not text:
        return None
    text = text.upper()
    if text in {"M", "MALE"}:
        return "M"
    if text in {"F", "FEMALE"}:
        return "F"
    return None


def extract_records():
    wb = openpyxl.load_workbook(RAW, data_only=True)
    rows = []
    for sheet_name, (year, month, month_name, sample_date) in MONTH_SHEETS.items():
        ws = wb[sheet_name]
        for excel_row in range(2, ws.max_row + 1):
            site = norm_site(ws.cell(excel_row, 3).value)
            sex = norm_sex(ws.cell(excel_row, 4).value)
            cl = as_float(ws.cell(excel_row, 5).value)
            if site not in {"DIAZ", "SWB"} or sex not in {"M", "F"} or cl is None:
                continue

            if sheet_name == "Dec 2024":
                profundidade_m = as_float(ws.cell(excel_row, 12).value)
                oxigenio_saturacao_pct = as_float(ws.cell(excel_row, 13).value)
                armadilha = as_int(ws.cell(excel_row, 14).value)
            else:
                profundidade_m = None
                oxigenio_saturacao_pct = as_float(ws.cell(excel_row, 12).value)
                armadilha = as_int(ws.cell(excel_row, 13).value)

            rows.append(
                {
                    "ano": year,
                    "mes": month,
                    "mes_nome": month_name,
                    "data_amostragem": sample_date.isoformat(),
                    "sitio": site,
                    "sexo": sex,
                    "comprimento_cefalotorax_mm": round(cl, 2),
                    "armadilha": armadilha,
                    "substrato": norm_text(ws.cell(excel_row, 8).value),
                    "temperatura_c": as_float(ws.cell(excel_row, 7).value),
                    "salinidade": as_float(ws.cell(excel_row, 9).value),
                    "oxigenio_dissolvido_mg_l": as_float(ws.cell(excel_row, 10).value),
                    "oxigenio_saturacao_pct": oxigenio_saturacao_pct,
                    "ph": as_float(ws.cell(excel_row, 11).value),
                    "profundidade_m": profundidade_m,
                    "coordenadas": norm_text(ws.cell(excel_row, 6).value),
                    "maturidade_reprodutiva": norm_text(ws.cell(excel_row, 14).value),
                }
            )

    for i, row in enumerate(rows, start=1):
        row["id_lagosta"] = f"LN{i:04d}"
    ordered = ["id_lagosta"] + [k for k in rows[0].keys() if k != "id_lagosta"]
    return [{k: row.get(k) for k in ordered} for row in rows]


def combined_counts():
    wb = openpyxl.load_workbook(RAW, data_only=True)
    ws = wb["Combined"]
    expected = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        month, site, total, males, females = row[:5]
        if month and site:
            expected[(COMBINED_MONTH[str(month)], norm_site(site))] = {
                "total": int(total),
                "M": int(males),
                "F": int(females),
            }
    return expected


def validate(rows):
    expected = combined_counts()
    observed = defaultdict(Counter)
    for row in rows:
        observed[(row["mes_nome"], row["sitio"])][row["sexo"]] += 1

    checks = []
    all_totals_ok = True
    all_sex_counts_ok = True
    for key, exp in sorted(expected.items()):
        obs_m = observed[key]["M"]
        obs_f = observed[key]["F"]
        obs_total = obs_m + obs_f
        total_ok = obs_total == exp["total"]
        sex_counts_ok = obs_m == exp["M"] and obs_f == exp["F"]
        all_totals_ok = all_totals_ok and total_ok
        all_sex_counts_ok = all_sex_counts_ok and sex_counts_ok
        checks.append(
            {
                "mes_nome": key[0],
                "sitio": key[1],
                "total_esperado": exp["total"],
                "total_observado": obs_total,
                "machos_esperado": exp["M"],
                "machos_observado": obs_m,
                "femeas_esperado": exp["F"],
                "femeas_observado": obs_f,
                "total_ok": total_ok,
                "contagens_sexo_ok": sex_counts_ok,
            }
        )

    missing = {
        key: sum(1 for row in rows if row[key] is None)
        for key in rows[0]
    }
    cl_values = [row["comprimento_cefalotorax_mm"] for row in rows]
    summary = {
        "linhas": len(rows),
        "colunas": len(rows[0]),
        "validacao_totais_ok": all_totals_ok,
        "validacao_contagens_sexo_ok": all_sex_counts_ok,
        "contagens": checks,
        "ausentes_por_coluna": missing,
        "comprimento_min_mm": min(cl_values),
        "comprimento_max_mm": max(cl_values),
        "comprimento_media_mm": round(mean(cl_values), 3),
        "duplicatas_completas": len(rows) - len({tuple(row.items()) for row in rows}),
    }
    VALIDATION_JSON.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def write_clean_csv(rows):
    with OUT_CSV.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def remove_sheet_if_exists(wb, sheet_name):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]


def set_header_comments(ws, comments):
    for cell in ws[1]:
        if cell.value in comments:
            cell.comment = Comment(comments[cell.value], "EAPA")


def autosize(ws, max_width=42):
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        values = [ws.cell(row, col_idx).value for row in range(1, min(ws.max_row, 80) + 1)]
        width = min(max_width, max(10, max(len(str(v)) if v is not None else 0 for v in values) + 2))
        ws.column_dimensions[letter].width = width


def update_workbook(rows, summary):
    wb = openpyxl.load_workbook(MAIN)
    remove_sheet_if_exists(wb, CLEAN_SHEET)
    ws = wb.create_sheet(CLEAN_SHEET)
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row[h] for h in headers])

    header_fill = PatternFill("solid", fgColor="1F6F78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)

    set_header_comments(
        ws,
        {
            "id_lagosta": "Identificador sequencial criado na curadoria; uma linha representa uma lagosta amostrada.",
            "ano": "Ano da amostragem.",
            "mes": "Mês da amostragem em formato numérico.",
            "mes_nome": "Mês da amostragem em português, sem acento para compatibilidade com R.",
            "data_amostragem": "Data de recuperação/amostragem inferida da aba mensal original.",
            "sitio": "Sítio de coleta indicado na fonte original: DIAZ ou SWB.",
            "sexo": "Sexo determinado no campo: M = macho; F = fêmea.",
            "comprimento_cefalotorax_mm": "Comprimento do cefalotórax da lagosta, em milímetros.",
            "armadilha": "Número da armadilha comercial modificada usada no evento de amostragem.",
            "substrato": "Descrição do substrato/habitat registrada na planilha original.",
            "temperatura_c": "Temperatura da água, em graus Celsius.",
            "salinidade": "Salinidade registrada no evento de amostragem.",
            "oxigenio_dissolvido_mg_l": "Oxigênio dissolvido, registrado como DO/DO2 em ppm na fonte; tratado aqui como mg/L.",
            "oxigenio_saturacao_pct": "Saturação de oxigênio dissolvido, em porcentagem.",
            "ph": "pH da água.",
            "profundidade_m": "Profundidade média, em metros, quando registrada; disponível no bruto apenas para dezembro.",
            "coordenadas": "Coordenadas textuais conforme a fonte original, sem tentativa de conversão geográfica.",
            "maturidade_reprodutiva": "Anotação reprodutiva ou observação associada ao indivíduo, quando registrada.",
        },
    )

    info = wb["info_conjuntos"]
    existing_ids = {info.cell(row, 1).value: row for row in range(2, info.max_row + 1)}
    row_idx = existing_ids.get(DATASET_ID, info.max_row + 1)
    info_values = {
        "id_conjunto": DATASET_ID,
        "nome_dataset_pacote": "lagosta_namibia",
        "nome_curto": "Lagosta-vermelha da costa oeste em florestas de kelp na Namibia",
        "objetivo_didatico": "Comparar abundancia, razao sexual e estrutura de tamanho de Jasus lalandii entre sitios e meses de amostragem.",
        "area": "bioecologia pesqueira; pesca costeira",
        "tema": "lagosta; estrutura de tamanho; razao sexual; habitat; qualidade da agua",
        "analise_principal": "ANOVA/GLM e tabelas de contingencia",
        "modelo_sugerido": "comprimento_cefalotorax_mm ~ sitio * sexo + mes_nome",
        "modelo_sugerido_r": "comprimento_cefalotorax_mm ~ sitio * sexo + mes_nome",
        "unidade_observacional": "lagosta individual amostrada por armadilha",
        "status_limpeza": "limpo_validado",
        "tipo_adaptacao": "abas mensais consolidadas em formato tidy; abas de normalidade/graficos e testes preliminares excluidas",
        "aba_limpa": CLEAN_SHEET,
        "linhas_limpo": summary["linhas"],
        "colunas_limpo": summary["colunas"],
        "linhas_dados_originais": "set/out/dez/jan/mar/abr: 1570 linhas brutas nas abas mensais",
        "colunas_dados_originais": "21 a 31 por aba mensal; dados analiticos concentrados nas 13 primeiras colunas",
        "fonte": "Malyenge, S. (2025). Abundance, Size Structure, and Sex Ratio Variation of West Coast Rock Lobster (Jasus lalandii) in Cultivated and Natural Kelp Forests off Lüderitz, Namibia. Mendeley Data.",
        "doi": "10.17632/y4vc857g9y.1",
        "url_documentacao": "https://data.mendeley.com/datasets/y4vc857g9y/1",
        "licenca": "CC BY 4.0",
        "uso_recomendado": "aula_livro_pacote_avaliacao",
        "arquivo_bruto_local": "provisorios/malyenge_2025_lobster_mendeley/Lobster Log sheet.xlsx",
        "observacoes": "Totais por mes/sitio conferidos contra a aba Combined. Ha uma divergencia interna na fonte em setembro/SWB: Combined informa total 167, mas o resumo de sexo soma 166; a base individual contem 101 machos e 66 femeas. Datas de dezembro foram corrigidas por contexto da aba mensal, pois o Excel interpretou 6/12/2024 como 12 de junho em algumas celulas.",
    }
    headers = [info.cell(1, col).value for col in range(1, info.max_column + 1)]
    for col_idx, header in enumerate(headers, start=1):
        info.cell(row_idx, col_idx).value = info_values.get(header)
    autosize(info)

    wb.save(MAIN)


def update_registry(summary):
    registry_row = {
        "id_fonte": "malyenge_2025_lagosta_namibia_mendeley",
        "origem": "Malyenge, S. (2025). Abundance, Size Structure, and Sex Ratio Variation of West Coast Rock Lobster (Jasus lalandii) in Cultivated and Natural Kelp Forests off Lüderitz, Namibia. Mendeley Data. DOI: 10.17632/y4vc857g9y.1.",
        "tipo_fonte": "dataset_mendeley",
        "data_acesso": "2026-07-11",
        "licenca": "CC BY 4.0",
        "status_licenca": "publico_aberto",
        "area": "bioecologia pesqueira; pesca costeira",
        "tema": "lagosta; estrutura de tamanho; razao sexual; habitat; qualidade da agua",
        "unidade_observacional": "lagosta individual",
        "estrutura_estatistica": "ANOVA/GLM; qui-quadrado; regressao logistica; descritiva por mes e sitio",
        "destino_pretendido": "aula_livro_pacote_avaliacao",
        "decisao": "aceitar",
        "arquivo_local": "dados_organizados_limpos_para_pacote_aulas.xlsx; provisorios/malyenge_2025_lobster_mendeley/Lobster Log sheet.xlsx",
        "observacoes": f"Base limpa adicionada na aba {CLEAN_SHEET} com {summary['linhas']} individuos e {summary['colunas']} variaveis. Totais por mes/sitio validados contra a aba Combined; divergencia interna de sexo em setembro/SWB documentada.",
    }
    with REGISTRY.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        rows = [row for row in reader if row.get("id_fonte") != registry_row["id_fonte"]]
        fieldnames = reader.fieldnames
    rows.append(registry_row)
    with REGISTRY.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main():
    rows = extract_records()
    summary = validate(rows)
    if not summary["validacao_totais_ok"]:
        raise SystemExit("Falha na validacao de contagens; conferir validacao_lagosta_namibia.json")
    write_clean_csv(rows)
    update_workbook(rows, summary)
    update_registry(summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
