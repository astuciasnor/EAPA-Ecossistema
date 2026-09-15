"""Monta o arquivo do aluno da atividade de regressao multipla (mercurio).

Rodar da raiz de ATIVIDADES/:

    python provisorios/mercurio_great_slave_lake/artifact/gerar_arquivo_aluno.py

Fonte: aba `Fish` do deposito Dryad de Rohonczy/Chételat et al. (2020), CC0.
258 peixes, uma linha por peixe. Cabecalho de duas linhas (8 e 9) achatado em uma;
`-` vira celula vazia; os 28 valores de Hg do musculo que os autores marcaram em
vermelho (estimados do corpo inteiro) ganham a coluna `Hg_muscle_estimado`.
"""

import datetime
import os
from collections import Counter

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ORIGEM = ("provisorios/mercurio_great_slave_lake/"
          "Rohonczy_et_al_2019_Journal_of_Great_Lakes_Research_mercury_"
          "Great_Slave_Lake_food_web_Nov_2019.xlsx")
SAIDA = "dados/mercurio_lago_escravo.xlsx"
NAVY = "FF0F3B5F"

COLS = ["Species", "Lab_ID", "Fish_ID", "Collection_Date", "Field_Site_ID",
        "Length_mm", "Weight_g", "Age_years", "Hg_muscle_mg_g_dw",
        "Hg_liver_mg_g_dw", "Hg_wholebody_mg_g_dw", "d13C_vpdb", "d15N_air"]
NUMERICAS = {"Length_mm", "Weight_g", "Age_years", "Hg_muscle_mg_g_dw",
             "Hg_liver_mg_g_dw", "Hg_wholebody_mg_g_dw", "d13C_vpdb", "d15N_air"}

ESPECIES = [
    ("BURB", "burbot"), ("CISC", "cisco"), ("INNC", "inconnu"),
    ("LKTR", "lake trout"), ("LKWH", "lake whitefish"),
    ("LNSC", "longnose sucker"), ("NRPK", "northern pike"),
    ("RNWH", "round whitefish"), ("WS", "white sucker"), ("SH", "shiner"),
    ("SLSC", "sculpin"), ("NSST", "ninespine stickleback"), ("WALL", "walleye"),
]

ORIGEM_SHEET = [
    ("Campo", "Informação"),
    ("Conjunto", "Mercúrio, idade, tamanho e isótopos em peixes do Great Slave Lake"),
    ("Fonte", "Rohonczy, J.; Cott, P. A.; Benwell, A.; Forbes, M. R.; Robinson, S. A.; "
              "Rosabal, M.; Amyot, M.; Chételat, J. (2020). Trophic structure and "
              "mercury transfer in the subarctic fish community of Great Slave Lake, "
              "Northwest Territories, Canada [Dataset]. Dryad. "
              "DOI: 10.5061/dryad.59zw3r23g (versão 5)"),
    ("Página", "https://datadryad.org/dataset/doi:10.5061/dryad.59zw3r23g"),
    ("Licença", "CC0 1.0 Universal (domínio público) — "
                "https://creativecommons.org/publicdomain/zero/1.0/"),
    ("Como citar", "Rohonczy et al. (2020), Dryad, DOI 10.5061/dryad.59zw3r23g. "
                   "Artigo: Journal of Great Lakes Research 46(2):402-413, "
                   "DOI 10.1016/j.jglr.2019.12.009"),
    ("Arquivo de origem", "aba `Fish` da planilha do depósito; cabeçalho em duas "
                          "linhas achatado em uma; `-` virou célula vazia"),
    ("Unidade observacional", "Peixe individual (coluna Fish_ID). 258 peixes de 13 "
                              "espécies, capturados em 2013-2015 na baía de "
                              "Yellowknife e no corpo principal do lago."),
    ("Variáveis", "Species (código; ver aba `especies`), Lab_ID, Fish_ID, "
                  "Collection_Date, Field_Site_ID, Length_mm, Weight_g, Age_years, "
                  "Hg_muscle_mg_g_dw, Hg_liver_mg_g_dw, Hg_wholebody_mg_g_dw, "
                  "d13C_vpdb, d15N_air (posição trófica)"),
    ("Atenção", "A coluna `Hg_muscle_estimado` marca SIM para os 28 peixes cuja "
                "concentração no músculo foi ESTIMADA a partir do corpo inteiro "
                "(os autores marcaram esses valores em vermelho na planilha "
                "original). Não são medidas diretas de músculo."),
    ("Ausentes", "Idade falta em 55 peixes; comprimento e peso em 12; Hg do músculo "
                 "em 2; δ15N em 1; Hg do fígado em 61 e do corpo inteiro em 222. "
                 "No arquivo original, ausente é o traço `-`."),
]

wb = Workbook()
ws = wb.active
ws.title = "dados"

fonte = load_workbook(ORIGEM, data_only=True)["Fish"]
estilo = load_workbook(ORIGEM)["Fish"]          # para ler a cor da fonte
ws.append(COLS + ["Hg_muscle_estimado"])

n_estimados = 0
for r in range(10, fonte.max_row + 1):
    v = [fonte.cell(row=r, column=c).value for c in range(1, 14)]
    if not any(x is not None and str(x).strip() for x in v):
        continue
    linha = []
    for nome, x in zip(COLS, v):
        if x is None or str(x).strip() in ("", "-", "NA"):
            linha.append(None)
        elif nome in NUMERICAS:
            linha.append(float(x))
        elif isinstance(x, (datetime.datetime, datetime.date)):
            linha.append(str(x)[:10])
        else:
            linha.append(str(x).strip())
    cor = estilo.cell(row=r, column=9).font.color
    rgb = getattr(cor, "rgb", None) if cor is not None else None
    estimado = isinstance(rgb, str) and rgb.upper() not in ("FF000000", "00000000")
    n_estimados += bool(estimado)
    linha.append("SIM" if estimado else None)
    ws.append(linha)

for celula in ws[1]:
    celula.font = Font(name="Calibri", size=10, bold=True, color="FFFFFFFF")
    celula.fill = PatternFill("solid", fgColor=NAVY)
    celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.freeze_panes = "A2"
for col, largura in zip("ABCDEFGHIJKLMN", [9, 12, 10, 15, 13, 12, 11, 11, 18, 17, 20, 11, 10, 18]):
    ws.column_dimensions[col].width = largura

ws2 = wb.create_sheet("especies")
ws2.append(["codigo", "nome_comum"])
for linha in ESPECIES:
    ws2.append(list(linha))
ws2.column_dimensions["A"].width = 10
ws2.column_dimensions["B"].width = 26

ws3 = wb.create_sheet("origem")
for linha in ORIGEM_SHEET:
    ws3.append(list(linha))
ws3.column_dimensions["A"].width = 22
ws3.column_dimensions["B"].width = 96
for celula in ws3[1]:
    celula.font = Font(name="Calibri", size=10, bold=True, color="FFFFFFFF")
    celula.fill = PatternFill("solid", fgColor=NAVY)
for linha in ws3.iter_rows(min_row=2):
    linha[1].alignment = Alignment(vertical="top", wrap_text=True)

os.makedirs("dados", exist_ok=True)
wb.save(SAIDA)
print(f"gravado: {SAIDA}")
print(f"peixes: {ws.max_row - 1} | colunas: {ws.max_column}")
print("Hg do músculo estimado (marcados SIM):", n_estimados)
print("abas:", wb.sheetnames)
