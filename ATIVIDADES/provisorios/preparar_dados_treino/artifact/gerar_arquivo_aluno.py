"""Monta o arquivo unico para o aluno praticar o menu Preparar Dados.

Rodar da raiz de ATIVIDADES/:

    python provisorios/preparar_dados_treino/artifact/gerar_arquivo_aluno.py

Junta, num unico xlsx, as abas de treino que hoje vivem espalhadas no repositorio
da IDE:

    CATALYSER/inst/app/dados/Treino-Transformacoes.xlsx
        biometria            -> aba `biometria`            (Trilha de Preparo, Calcular, Contingencia)
        desembarques_largo   -> aba `desembarques_largo`   (Arrumar > empilhar)
        guia                 -> NAO entra: o mapa vai para o README.md do candidato
    CATALYSER/inst/app/dados/Treino-Arrumacao.xlsx
        Treino-Largo         -> aba `receitas_largo`       (Arrumar > empilhar, mais rica)
        Treino-Separar       -> aba `coletas_composta`     (Arrumar > separar)

As imperfeicoes sao copiadas **como estao** (caixa e espacos inconsistentes, NA,
duplicatas, formato largo, coluna composta): limpar e o exercicio do aluno. So o
cabecalho recebe formatacao.

A aba `guia` do arquivo original **nao entra**: ela diz em qual coluna esta cada
problema e funcionaria como gabarito. O mapa vive em
`provisorios/preparar_dados_treino/README.md`, fora do arquivo do aluno. A aba
`origem`, sim, e escrita por este script.
"""

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_IDE = r"D:\Claude\EAPA-Ecossistema\CATALYSER\inst\app\dados"
TRANSFORMACOES = f"{BASE_IDE}\\Treino-Transformacoes.xlsx"
ARRUMACAO = f"{BASE_IDE}\\Treino-Arrumacao.xlsx"
SAIDA = "dados/preparar_dados_treino.xlsx"

NAVY = "FF0F3B5F"

# (arquivo de origem, aba de origem, nome da aba na saida, larguras por coluna)
ABAS = [
    (TRANSFORMACOES, "biometria", "biometria",
     [6, 16, 14, 9, 20, 13, 15, 10, 8, 14]),
    (TRANSFORMACOES, "desembarques_largo", "desembarques_largo",
     [16, 16, 16, 16]),
    (ARRUMACAO, "Treino-Largo", "receitas_largo",
     [14, 12, 15, 16, 15, 16, 15, 16]),
    (ARRUMACAO, "Treino-Separar", "coletas_composta",
     [6, 18, 16, 15, 10]),
]

ORIGEM = [
    ("Campo", "Informação"),
    ("Arquivo", "preparar_dados_treino.xlsx"),
    ("O que é", "Arquivo único de treino para o aluno percorrer o menu Preparar Dados "
                "da CatalyseR: Trilha de Preparo, Calcular/Reescalar, Arrumar e Contingência."),
    ("De onde vem", "Duas planilhas de treino do repositório da IDE: "
                    "CATALYSER/inst/app/dados/Treino-Transformacoes.xlsx (abas biometria, "
                    "desembarques_largo, guia) e Treino-Arrumacao.xlsx (abas Treino-Largo, "
                    "Treino-Separar). Os valores foram copiados sem alteração."),
    ("Natureza dos dados", "SINTÉTICOS, de propósito. Dado real tem problemas, mas "
                           "raramente todos ao mesmo tempo e reconhecíveis: aqui cada "
                           "imperfeição foi plantada para exercitar um tratamento "
                           "específico do menu. Não são observações de campo nem servem "
                           "para conclusão biológica ou pesqueira."),
    ("Licença", "Não se aplica: material interno da disciplina, sem fonte externa. "
                "Por isso o conjunto não entra no EAPADados."),
    ("Para que serve", "Praticar o PREPARO. As atividades AVALIADAS do pilar usam dado "
                       "real e distinto deste; este arquivo é o aquecimento do menu. "
                       "Uso temporário em ATIVIDADES/dados/."),
    ("O que está plantado", "Caixa e espaços inconsistentes em local, sexo e especie; "
                            "valores ausentes em comprimento_cm e peso_g; 3 linhas "
                            "duplicadas; formato largo em duas abas; colunas compostas "
                            "em coletas_composta."),
    ("Mapa do preparo", "Não está neste arquivo de propósito: o mapa (o que exercitar e "
                        "em qual coluna) fica em provisorios/preparar_dados_treino/README.md, "
                        "como gabarito do professor. O aluno recebe só os dados."),
    ("Unidade observacional", "biometria e coletas_composta: um peixe por linha. "
                              "desembarques_largo e receitas_largo: um porto (ou porto x "
                              "espécie) por linha, com anos nas colunas."),
]


def copiar_aba(wb_saida, caminho, aba_origem, nome_saida, larguras):
    ws_origem = load_workbook(caminho, data_only=True)[aba_origem]
    linhas = [l for l in ws_origem.iter_rows(values_only=True)
              if any(v is not None and str(v).strip() for v in l)]
    ws = wb_saida.create_sheet(nome_saida)
    for linha in linhas:
        ws.append(list(linha))
    for col, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(col)].width = largura
    for celula in ws[1]:
        celula.font = Font(name="Calibri", size=10, bold=True, color="FFFFFFFF")
        celula.fill = PatternFill("solid", fgColor=NAVY)
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    print(f"  {nome_saida}: {ws.max_row - 1} linhas x {ws.max_column} colunas "
          f"(de {caminho.split(chr(92))[-1]} :: {aba_origem})")
    return ws


def escrever_tabela(wb, nome, dados, larguras, titulo=None):
    ws = wb.create_sheet(nome)
    if titulo:
        ws.append([titulo])
        celula = ws.cell(row=1, column=1)
        celula.font = Font(name="Calibri", size=11, bold=True, color=NAVY)
        ws.append([])
    for linha in dados:
        ws.append(list(linha))
    linha_cab = 3 if titulo else 1
    for celula in ws[linha_cab]:
        celula.font = Font(name="Calibri", size=10, bold=True, color="FFFFFFFF")
        celula.fill = PatternFill("solid", fgColor=NAVY)
        celula.alignment = Alignment(horizontal="left", vertical="center")
    for col, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(col)].width = largura
    ws.freeze_panes = f"A{linha_cab + 1}"
    return ws


wb = Workbook()
wb.remove(wb.active)

print("copiando abas de treino:")
for caminho, aba, nome, larguras in ABAS:
    copiar_aba(wb, caminho, aba, nome, larguras)

escrever_tabela(wb, "origem", ORIGEM, [30, 96])

wb.save(SAIDA)
print(f"\ngravado: {SAIDA}")
print(f"abas: {wb.sheetnames}")
